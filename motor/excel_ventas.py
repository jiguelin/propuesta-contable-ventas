# -*- coding: utf-8 -*-
"""
motor/excel_ventas.py — Excel de importación para NewContaSis (FORMATO_VENTAS, 44 columnas A–AR),
Excel de revisión y reporte de texto.

El Excel de importación lleva SOLO filas de datos desde la fila 1 (la nota de la plantilla exige
borrar las filas 1 a 13). Una fila por comprobante.
"""
from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .clasificador_ventas import ConfigEmpresa, Fila

D2 = Decimal('0.01')

# (letra, cabecera, tipo)  tipo: f = fecha, t = texto, n = número 2 dec., tc = tipo de cambio
COLUMNAS = [
    ('A', 'FECHA EMISION', 'f'), ('B', 'FECHA VENCIMIENTO O PAGO', 'f'), ('C', 'TIPO CP', 't'),
    ('D', 'SERIE', 't'), ('E', 'NUMERO', 't'), ('F', 'TIPO DOC CLIENTE', 't'), ('G', 'NUMERO DOC CLIENTE', 't'),
    ('H', 'CLIENTE', 't'), ('I', 'VALOR EXPORTACION', 'n'), ('J', 'BASE IMPONIBLE', 'n'), ('K', 'EXONERADA', 'n'),
    ('L', 'INAFECTA', 'n'), ('M', 'ISC', 'n'), ('N', 'IGV', 'n'), ('O', 'OTROS TRIBUTOS', 'n'), ('P', 'TOTAL', 'n'),
    ('Q', 'TIPO DE CAMBIO', 'tc'), ('R', 'FECHA DOC MODIFICADO', 'f'), ('S', 'TIPO DOC MODIFICADO', 't'),
    ('T', 'SERIE DOC MODIFICADO', 't'), ('U', 'NUMERO DOC MODIFICADO', 't'), ('V', 'MONEDA', 't'),
    ('W', 'EQUIVALENTE EN DOLARES', 'n'), ('X', 'FECHA VENCIMIENTO', 'f'), ('Y', 'CONDICION', 't'),
    ('Z', 'CENTRO COSTOS', 't'), ('AA', 'CENTRO COSTOS 2', 't'), ('AB', 'CUENTA BASE IMPONIBLE', 't'),
    ('AC', 'CUENTA OTROS TRIBUTOS', 't'), ('AD', 'CUENTA TOTAL', 't'), ('AE', 'REGIMEN ESPECIAL', 't'),
    ('AF', '% REGIMEN', 'n'), ('AG', 'IMPORTE REGIMEN', 'n'), ('AH', 'SERIE DOC REGIMEN', 't'),
    ('AI', 'NUMERO DOC REGIMEN', 't'), ('AJ', 'FECHA DOC REGIMEN', 'f'), ('AK', 'CODIGO PRESUPUESTO', 't'),
    ('AL', '% IGV', 'n'), ('AM', 'GLOSA', 't'), ('AN', 'MEDIO DE PAGO', 't'), ('AO', 'CONDICION PERCEPCION', 't'),
    ('AP', 'IMPORTE BASE REGIMEN', 'n'), ('AQ', 'ICBPER', 'n'), ('AR', 'CUENTA ICBPER', 't'),
]
assert len(COLUMNAS) == 44


def _fecha(iso: str):
    try:
        return datetime.strptime((iso or '')[:10], '%Y-%m-%d')
    except ValueError:
        return None


def _q(v) -> Decimal:
    return Decimal(v).quantize(D2, ROUND_HALF_UP)


def importes(f: Fila, cfg: ConfigEmpresa) -> dict:
    """I..P, AQ y W en la moneda de registro (soles si usd_en_soles), con signo de NC y cuadre de centavos."""
    cero = {k: Decimal('0.00') for k in ('I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'AQ')}
    if f.c is None or f.estado in ('anulado', 'hueco_anulado'):
        return {**cero, 'W': None}
    c = f.c
    signo = Decimal(-1) if (f.tipo == '07' and cfg.nc_negativo) else Decimal(1)
    usd = c.moneda_codigo == 'USD'
    k = Decimal(str(f.tc)) if (usd and cfg.usd_en_soles and f.tc) else Decimal(1)
    v = {
        'I': _q(c.exportacion * k), 'J': _q(c.gravadas * k), 'K': _q(c.exoneradas * k), 'L': _q(c.inafectas * k),
        'M': _q(c.isc * k), 'N': _q(c.igv * k), 'O': _q((c.otros_tributos + c.otros_cargos) * k),
        'P': _q(c.importe_total * k), 'AQ': _q(c.icbper * k),
    }
    if k != 1:   # el redondeo de la conversión se cuadra contra el IGV (o la base si no hay IGV)
        dif = v['P'] - sum(v[x] for x in ('I', 'J', 'K', 'L', 'M', 'N', 'O', 'AQ'))
        if 0 < abs(dif) <= Decimal('0.05'):
            col = 'N' if v['N'] else ('J' if v['J'] else 'I')
            v[col] += dif
    v = {x: y * signo for x, y in v.items()}
    v['W'] = _q(c.importe_total) * signo if usd else None
    return v


def fila_excel(f: Fila, cfg: ConfigEmpresa) -> list:
    m = importes(f, cfg)
    otros = m['O'] != 0
    tasa = f.c.tasa_igv if (f.c is not None and f.c.igv) else None
    return [
        _fecha(f.fecha), _fecha(f.fecha_venc or f.fecha), f.tipo, f.serie, str(f.numero),           # A-E
        f.cli_tipo, f.cli_doc, (f.cli_nombre or '')[:60],                                           # F-H
        m['I'], m['J'], m['K'], m['L'], m['M'], m['N'], m['O'], m['P'],                              # I-P
        f.tc,                                                                                        # Q (siempre)
        _fecha(f.ref_fecha) if f.tipo in ('07', '08') else None,                                     # R
        f.ref_tipo if f.tipo in ('07', '08') else '', f.ref_serie if f.tipo in ('07', '08') else '',
        f.ref_numero if f.tipo in ('07', '08') else '',                                              # S-U
        f.moneda, m['W'], _fecha(f.fecha_venc or f.fecha), f.condicion,                              # V-Y
        '', '', f.cuenta_base, f.cuenta_otros if otros else '', f.cuenta_total,                      # Z-AD
        '', None, None, '', '', None, '',                                                            # AE-AK (régimen especial: no se registra en ventas)
        float(tasa) if tasa else None, (f.glosa or '')[:60],                                         # AL-AM
        f.medio_pago if f.condicion == 'CON' else '', '', None,                                      # AN-AP
        m['AQ'] if m['AQ'] else None, f.cuenta_icbper if m['AQ'] else '',                            # AQ-AR
    ]


def _escribir(ws, filas: list[list], inicio: int = 1):
    for i, fila in enumerate(filas):
        for j, (val, (_, _, tipo)) in enumerate(zip(fila, COLUMNAS), start=1):
            if isinstance(val, Decimal):
                val = float(val)
            cell = ws.cell(row=inicio + i, column=j, value=val if val != '' else None)
            if tipo == 'f':
                cell.number_format = 'dd/mm/yyyy'
            elif tipo == 'n':
                cell.number_format = '0.00'
            elif tipo == 'tc':
                cell.number_format = '0.000'
            else:
                cell.number_format = '@'
                if val is not None and val != '':
                    cell.value = str(val)


def generar_importacion(filas: list[Fila], cfg: ConfigEmpresa) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'FORMATO_VENTAS'
    _escribir(ws, [fila_excel(f, cfg) for f in filas if f.exportable])
    for j in range(1, 45):
        ws.column_dimensions[get_column_letter(j)].width = 14
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------- revisión
_CAB = Font(bold=True, color='FFFFFF')
_FONDO = PatternFill('solid', fgColor='1F3B5C')


def _hoja(wb, titulo: str, cab: list[str], filas: list[list], anchos: dict[int, int] | None = None):
    ws = wb.create_sheet(titulo)
    ws.append(cab)
    for c in ws[1]:
        c.font, c.fill, c.alignment = _CAB, _FONDO, Alignment(wrap_text=True, vertical='center')
    for r in filas:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in r])
    ws.freeze_panes = 'A2'
    for j in range(1, len(cab) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (anchos or {}).get(j, 14)
    return ws


def generar_revision(lote, cfg: ConfigEmpresa, plan=None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    desc = (lambda cod: (plan.get(cod).descripcion if plan and plan.get(cod) else '')) if plan else (lambda cod: '')
    cab = ['Semáforo', 'Estado', 'Fecha', 'Tipo', 'Serie-Número', 'Tipo doc', 'Doc cliente', 'Cliente', 'Moneda',
           'Base', 'IGV', 'Total', 'TC', 'Condición', 'Cuenta base (AB)', 'Descripción AB', 'Cuenta total (AD)',
           'Concepto', 'Confianza', 'Origen', 'Por qué', 'Glosa', 'Motivo', 'Avisos']

    def fila_rev(f: Fila):
        c = f.c
        return [f.semaforo, f.estado, f.fecha, f.tipo, f.id, f.cli_tipo, f.cli_doc, f.cli_nombre or (c.nombre_cliente if c else ''),
                f.moneda, c.gravadas if c else None, c.igv if c else None, c.importe_total if c else None, f.tc,
                f.condicion, f.cuenta_base, desc(f.cuenta_base), f.cuenta_total, f.concepto, f.confianza, f.origen,
                f.explicacion, f.glosa, f.motivo, '; '.join(f.avisos)]
    anchos = {5: 14, 8: 34, 16: 34, 21: 40, 22: 34, 23: 40, 24: 60}
    _hoja(wb, 'PROPUESTA', cab, [fila_rev(f) for f in lote.filas if f.estado in ('ok', 'anulado')], anchos)
    _hoja(wb, 'REVISAR', cab, [fila_rev(f) for f in lote.filas if f.estado == 'revisar'], anchos)
    _hoja(wb, 'ANULADOS Y FALTANTES', cab,
          [fila_rev(f) for f in lote.filas if f.estado in ('anulado', 'hueco_anulado', 'faltante')], anchos)
    _hoja(wb, 'EXCLUIDOS', ['Archivo / comprobante', 'Fecha', 'Motivo'],
          [[(f.c.archivo if f.c else f.id), f.fecha, f.motivo] for f in lote.filas if f.estado == 'excluido'] +
          [[n, '', m] for n, m in lote.no_procesados], {1: 40, 3: 70})
    _hoja(wb, 'ANTICIPOS', ['Comprobante final', 'Fecha', 'Anticipo', 'Monto anticipo', 'Base', 'Base en soles',
                            'Debe', 'Haber', 'Dónde se registró el anticipo', 'Acción'],
          [[a['final'], a['fecha'], a['anticipo'], a['monto_anticipo'], a['base'], a['base_soles'], a['debe'] if a['proponer'] else '',
            a['haber'] if a['proponer'] else '', a['origen'], a['accion']] for a in lote.anticipos], {9: 24, 10: 70})
    _hoja(wb, 'ENTIDADES NUEVAS', ['Código', 'Nombre', 'ID extranjero', 'País'],
          [[e['codigo'], e['nombre'], e.get('id_extranjero', ''), e.get('pais', '')] for e in lote.entidades_nuevas], {2: 50})
    ws = wb.create_sheet('NEWCONTASIS (ref)')
    ws.append([f'{l} {n}' for l, n, _ in COLUMNAS])
    for c in ws[1]:
        c.font, c.fill = _CAB, _FONDO
    _escribir(ws, [fila_excel(f, cfg) for f in lote.filas if f.exportable], inicio=2)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generar_reporte(lote, cfg: ConfigEmpresa) -> str:
    exp = lote.exportables
    tot = {k: Decimal(0) for k in ('J', 'N', 'P')}
    for f in exp:
        m = importes(f, cfg)
        for k in tot:
            tot[k] += m[k]
    lin = [f'REPORTE DE VENTAS — {cfg.nombre} (RUC {cfg.ruc}) — periodo {lote.periodo}',
           f'Generado: {datetime.now():%d/%m/%Y %H:%M}', '',
           f'Filas para importar: {len(exp)}   Base: {tot["J"]:,.2f}   IGV: {tot["N"]:,.2f}   Total: {tot["P"]:,.2f} (soles)',
           f'Contado: {sum(1 for f in exp if f.c is not None and f.condicion == "CON")}   '
           f'Crédito: {sum(1 for f in exp if f.c is not None and f.condicion == "CRE")}',
           f'Por revisar: {len(lote.por_estado("revisar"))}   Excluidos: {len(lote.por_estado("excluido"))}   '
           f'Anulados: {len(lote.por_estado("anulado"))}   Huecos propuestos como anulados: {len(lote.por_estado("hueco_anulado"))}   '
           f'XML faltantes: {len(lote.por_estado("faltante"))}', '']
    if lote.bloqueos:
        lin += ['BLOQUEOS (no se puede importar):'] + [f'  - {b}' for b in lote.bloqueos] + ['']
    if cfg.modo_contado == 'directo' and any(f.condicion == 'CON' for f in exp):
        lin += ['IMPORTANTE: hay ventas al contado con la cuenta de caja en AD. Al importar en NewContaSis marque',
                '"Generar solo - Asientos de Venta" para que no se genere un segundo asiento de cobro.', '']
    secciones = [('POR REVISAR', lote.por_estado('revisar'), lambda f: f.motivo),
                 ('HUECOS EN LA NUMERACIÓN (propuestos como anulados)', lote.por_estado('hueco_anulado'), lambda f: 'confirmado' if f.confirmado else 'SIN CONFIRMAR: no se exporta'),
                 ('XML FALTANTES', lote.por_estado('faltante'), lambda f: f.motivo),
                 ('ANULADOS', lote.por_estado('anulado'), lambda f: f.motivo),
                 ('EXCLUIDOS', lote.por_estado('excluido'), lambda f: f.motivo)]
    for titulo, filas, txt in secciones:
        if filas:
            lin.append(f'{titulo} ({len(filas)}):')
            lin += [f'  {f.tipo} {f.id:14} {f.fecha}  {txt(f)}' for f in filas]
            lin.append('')
    avisos = [(f, a) for f in exp + lote.por_estado('revisar') for a in f.avisos]
    if avisos:
        lin.append(f'AVISOS ({len(avisos)}):')
        lin += [f'  {f.tipo} {f.id:14} {a}' for f, a in avisos]
        lin.append('')
    prop = [a for a in lote.anticipos if a['proponer']]
    if prop:
        lin.append('ASIENTOS COMPLEMENTARIOS DE ANTICIPOS (registrar a mano en NewContaSis):')
        for a in prop:
            lin.append(f"  {a['final']:14} aplica {a['anticipo']:12} Debe {a['debe']} / Haber {a['haber']}  {a['base_soles']:,.2f}  ({a['origen']})")
        lin.append('')
    if lote.entidades_nuevas:
        lin.append('CLIENTES DEL EXTERIOR NUEVOS (se crean al importar; revíselos en Entidades):')
        lin += [f"  {e['codigo']}  {e['nombre']}  {e.get('pais', '')}" for e in lote.entidades_nuevas]
        lin.append('')
    if lote.no_procesados:
        lin.append('ARCHIVOS NO PROCESADOS:')
        lin += [f'  {n}: {m}' for n, m in lote.no_procesados]
    return '\n'.join(lin)
