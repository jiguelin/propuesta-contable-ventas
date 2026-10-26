# -*- coding: utf-8 -*-
"""Casos de aceptación de la especificación (sección 13), con XML reales anonimizados y XML sintéticos."""
from decimal import Decimal

import io

import openpyxl

from conftest import ESTUDIO, JOYERIA, MINERA, archivos
from fabrica import baja, comprobante
from datetime import date

from motor.clasificador_ventas import ConfigEmpresa, cuenta_total
from motor.tipo_cambio import TC
from motor.excel_ventas import COLUMNAS, fila_excel, generar_importacion, importes
from motor.lector import leer_archivos
from motor.proceso import aplicar_cambio
from motor.xml_parser import corregir_codificacion, parse_xml


def fila(lote, sid):
    return next(f for f in lote.filas if f.id == sid)


def col(fila_xl, letra):
    return fila_xl[[c[0] for c in COLUMNAS].index(letra)]


# ----------------------------------------------------------------------------- lector
def test_los_58_xml_reales_se_leen_sin_error():
    arch = archivos('servicios_oct') + archivos('joyeria_set') + archivos('dolares_set')
    lec = leer_archivos(arch)
    assert len(lec.comprobantes) == 58
    assert not [c for c in lec.comprobantes if c.error]


def test_utf8_declarado_iso_se_corrige():                       # caso 26
    data = next(d for n, d in archivos('joyeria_set') if 'BON1-524' in n)
    assert b'ISO-8859-1' in data[:80]
    assert b'encoding="UTF-8"' in corregir_codificacion(data)[:80]
    assert parse_xml(data).nombre_cliente == 'ÑUÑEZ DÍAZ PRUEBA'


def test_vencimiento_sale_de_la_cuota():
    c = parse_xml(next(d for n, d in archivos('servicios_oct') if 'E001-1099' in n))
    assert c.fecha_vencimiento == '2026-10-30'


# ----------------------------------------------------------------------------- plan y mapa
def test_plan_hojas_por_prefijo(plan):
    assert plan.ruc == ESTUDIO and plan.giro == 'CONSULTORIA EMPRESARIAL'
    assert not plan.get('70321').es_hoja
    assert [c.codigo for c in plan.hojas_bajo('70321')] == ['7032101', '7032102', '7032103']


def test_mapa_sugerido(mapa):
    assert mapa['SERVICIO_LOCAL'].es_grupo and mapa['SERVICIO_LOCAL'].hojas == ['7032101', '7032102', '7032103']
    assert mapa['VENTA_MERCADERIA_LOCAL'].cuenta == '70121'
    assert mapa['VENTA_ACTIVO_FIJO'].cuenta == '7564'           # no existe 75641
    assert mapa['ANTICIPO_RECIBIDO'].cuenta == '1221'
    assert (mapa['CAJA_PEN'].cuenta, mapa['CAJA_USD'].cuenta, mapa['CXC_PEN'].cuenta) == ('1011', '1012', '1212')


# ----------------------------------------------------------------------------- servicios (CONTADEUS anonimizado)
def test_servicios_octubre_bloquea_sin_tc(correr):              # caso 18
    lote, _ = correr(ESTUDIO, 'servicios', archivos('servicios_oct'), '2026-10')
    assert any('tipo de cambio' in b for b in lote.bloqueos)


def test_servicios_cuentas_y_vencimiento(correr, tc):           # casos 21 y 22
    lote, ctx = correr(ESTUDIO, 'servicios', archivos('servicios_oct'), '2026-10',
                       tc_extra=[TC(date(2026, 10, 22), 3.390, 3.400, 'prueba')])
    f = fila(lote, 'E001-1099')
    assert (f.estado, f.condicion, f.cuenta_total, f.cuenta_base) == ('ok', 'CRE', '1212', '7032101')
    x = fila_excel(f, ctx.cfg)
    assert col(x, 'B').strftime('%d/%m/%Y') == col(x, 'X').strftime('%d/%m/%Y') == '30/10/2026'
    assert col(x, 'P') == Decimal('3722.90')
    assert all(col(x, k) in ('', None) for k in ('AE', 'AF', 'AG', 'AH', 'AI', 'AJ', 'AP'))   # detracción: solo informativa
    f2 = fila(lote, 'E001-1110')                                # "SERIVICIO CONTABLE" + "SERVICIO LABORAL"
    assert f2.cuenta_base == '7032101' and not f2.mixta
    assert fila(lote, 'E001-1111').cuenta_total == '1011'       # contado en soles
    assert not lote.por_estado('hueco_anulado')


def test_fuera_de_periodo(correr):                              # casos 7 y 23
    lote, _ = correr(ESTUDIO, 'servicios', archivos('servicios_oct'), '2026-09')
    assert len(lote.por_estado('excluido')) == 28
    assert all('Fuera de periodo' in f.motivo for f in lote.por_estado('excluido'))


# ----------------------------------------------------------------------------- joyería (ESSENTTA anonimizado)
def test_huecos_anulados_y_faltante(correr):                    # caso 24
    lote, _ = correr(JOYERIA, 'comercial', archivos('joyeria_set'))
    huecos = sorted(f.numero for f in lote.por_estado('hueco_anulado'))
    assert huecos == [4480, 4482, 4483, 4484, 4488, 4489, 4490, 4497]
    assert [f.id for f in lote.por_estado('faltante')] == ['BOP1-4494']
    h = fila(lote, 'BOP1-4480')
    assert (h.cli_tipo, h.cli_doc, h.cli_nombre) == ('0', '0000', 'ANULADO')
    assert not h.exportable                                     # sin confirmar no se exporta
    h.confirmado = True
    assert h.exportable


def test_anticipos(correr):                                     # caso 25
    lote, ctx = correr(JOYERIA, 'comercial', archivos('joyeria_set'))
    a = fila(lote, 'BOP1-4477')
    assert a.es_anticipo and a.cuenta_base == '1221' and a.glosa == 'ANTICIPO DE CLIENTE'
    f = fila(lote, 'BOP1-4495')
    m = importes(f, ctx.cfg)
    assert (m['J'], m['N'], m['P']) == (Decimal('1271.19'), Decimal('228.81'), Decimal('1500.00'))
    assert f.cuenta_base == '70121'
    de_4495 = [x for x in lote.anticipos if x['final'] == 'BOP1-4495']
    assert {x['anticipo'] for x in de_4495} == {'BOP1-4384', 'BOP1-4477', 'BOP1-4494'}
    assert sum(x['base'] for x in de_4495) == Decimal('7542.37')   # = descuento global 04 del XML
    assert all((x['debe'], x['haber']) == ('1221', '70121') for x in de_4495)
    assert len(lote.anticipos) == 10


def test_boletas_con_dni_y_carne(correr):
    lote, ctx = correr(JOYERIA, 'comercial', archivos('joyeria_set'))
    assert fila(lote, 'BOP1-4481').cli_tipo == '4' and fila(lote, 'BOP1-4481').cli_doc.startswith('000')
    assert all(f.cuenta_base in ('70121', '1221') for f in lote.por_estado('ok') if f.c is not None)
    assert all(f.cuenta_total == '1011' for f in lote.por_estado('ok') if f.c is not None)


# ----------------------------------------------------------------------------- dólares
def test_dolares_al_contado(correr):                            # caso 27
    lote, ctx = correr(MINERA, 'mixta', archivos('dolares_set'))
    f = fila(lote, 'E001-130')
    m = importes(f, ctx.cfg)
    assert f.tc == 3.368 and f.cuenta_total == '1012' and f.moneda == 'D'
    assert (m['J'], m['N'], m['P'], m['W']) == (Decimal('5052.00'), Decimal('909.36'), Decimal('5961.36'), Decimal('1770.00'))
    g = fila(lote, 'E001-131')
    mg = importes(g, ctx.cfg)
    assert mg['J'] + mg['N'] == mg['P']                         # cuadre tras convertir


# ----------------------------------------------------------------------------- boletas sin DNI
def test_boleta_sin_dni(correr):                                # casos 4 y 5
    arch = [('b1.xml', comprobante('B001-10', tipo='03', cli_doc='-', cli_scheme='0', cli_nombre='-',
                                   items=(('ANILLO PLATA', 16.95, 'NIU'),), ruc_emisor=JOYERIA)),
            ('b2.xml', comprobante('B001-11', tipo='03', cli_doc='00000000', cli_scheme='1', cli_nombre='CLIENTE VARIOS',
                                   items=(('ANILLO ORO', 800.0, 'NIU'),), ruc_emisor=JOYERIA))]
    lote, _ = correr(JOYERIA, 'comercial', arch)
    b1, b2 = fila(lote, 'B001-10'), fila(lote, 'B001-11')
    assert (b1.cli_tipo, b1.cli_doc, b1.cli_nombre) == ('0', '03', 'BOLETA DE VENTA') and not b1.avisos
    assert b2.cli_doc == '03' and any('700' in a for a in b2.avisos)


# ----------------------------------------------------------------------------- exportación y entidades
def test_exportacion_codigos_de_exterior(correr):               # casos 13 y 28
    arch = [('e1.xml', comprobante('E001-200', cli_doc='-', cli_scheme='0', cli_nombre='Globex Trading, LLC', pais='US',
                                   exportacion=True, operacion='0201', items=(('CONSULTING SERVICES', 1000.0, 'ZZ'),))),
            ('e2.xml', comprobante('E001-201', cli_doc='-', cli_scheme='0', cli_nombre='ACME WIDGETS INC', pais='US',
                                   exportacion=True, operacion='0201', items=(('CONSULTING SERVICES', 500.0, 'ZZ'),)))]
    lote, ctx = correr(ESTUDIO, 'servicios', arch)
    e1, e2 = fila(lote, 'E001-200'), fila(lote, 'E001-201')
    assert e1.es_exportacion and e1.cli_doc == '#0000000009'
    assert e1.cuenta_base == '70311'
    m = importes(e1, ctx.cfg)
    assert (m['I'], m['J'], m['N'], m['P']) == (Decimal('1000.00'), Decimal('0.00'), Decimal('0.00'), Decimal('1000.00'))
    assert e2.cli_doc == '10000000004'                          # 10000000001 y 10000000003 ya existen
    assert lote.entidades_nuevas[0]['codigo'] == '10000000004'
    # el mes siguiente el mismo cliente recibe el mismo código, sin volver a subir entidades
    lote2, _ = correr(ESTUDIO, 'servicios', [('e3.xml', comprobante('E001-300', fecha='2026-09-20', cli_doc='-', cli_scheme='0',
                                                                    cli_nombre='Acme Widgets, Inc.', pais='US', exportacion=True,
                                                                    operacion='0201'))],
                      registros_exterior=lote.entidades_nuevas)
    assert fila(lote2, 'E001-300').cli_doc == '10000000004' and not lote2.entidades_nuevas


# ----------------------------------------------------------------------------- notas de crédito
def test_nota_de_credito(correr):                               # caso 10
    arch = [('f.xml', comprobante('F001-50', fecha='2026-09-02', forma_pago='Credito',
                                  items=(('SERVICIO DE CAPACITACION', 500.0, 'ZZ'),))),
            ('nc.xml', comprobante('FC01-5', fecha='2026-09-15', tipo='07', nc_ref=('F001-50', '01', '2026-09-02'),
                                   items=(('SERVICIO DE CAPACITACION', 100.0, 'ZZ'),)))]
    lote, ctx = correr(ESTUDIO, 'servicios', arch)
    f, nc = fila(lote, 'F001-50'), fila(lote, 'FC01-5')
    assert f.cuenta_base == '7032103'                            # subcuenta de capacitaciones
    assert nc.cuenta_base == '7032103' and nc.condicion == 'CRE' and nc.cuenta_total == '1212'
    x = fila_excel(nc, ctx.cfg)
    assert (col(x, 'S'), col(x, 'T'), col(x, 'U')) == ('01', 'F001', '50')
    assert col(x, 'R').strftime('%d/%m/%Y') == '02/09/2026'
    assert col(x, 'P') == Decimal('-118.00')


def test_nota_sin_fecha_de_referencia_va_a_revisar(correr):
    arch = [('nc.xml', comprobante('FC01-6', tipo='07', nc_ref=('F001-999', '01', ''),
                                   items=(('SERVICIO DE CAPACITACION', 100.0, 'ZZ'),)))]
    lote, _ = correr(ESTUDIO, 'servicios', arch)
    assert fila(lote, 'FC01-6').estado == 'revisar'


# ----------------------------------------------------------------------------- bajas, duplicados, otra empresa
def test_comunicacion_de_baja_y_resumen(correr):                # caso 6
    arch = [('f1.xml', comprobante('F001-1')), ('f2.xml', comprobante('F001-2')), ('f3.xml', comprobante('F001-3')),
            ('ra.xml', baja([('01', 'F001', 2)])), ('rc.xml', baja([('03', 'B001', 7)], resumen=True))]
    lote, _ = correr(ESTUDIO, 'servicios', arch)
    an = fila(lote, 'F001-2')
    assert an.estado == 'anulado' and an.cli_doc == '0000' and an.exportable
    assert fila(lote, 'B001-7').estado == 'anulado'
    assert importes(an, ConfigEmpresa(ruc=ESTUDIO))['P'] == 0


def test_duplicado_y_otra_empresa(correr):                      # casos 8 y 9
    x = comprobante('F001-1')
    arch = [('a.xml', x), ('b.xml', x), ('c.xml', comprobante('F001-2', ruc_emisor='20999999999'))]
    lote, _ = correr(ESTUDIO, 'servicios', arch)
    motivos = sorted(f.motivo[:9] for f in lote.por_estado('excluido'))
    assert motivos == ['Duplicado', 'El emisor']


def test_cdr_y_pdf_no_se_procesan():
    lec = leer_archivos([('R-1.xml', b'<?xml version="1.0"?><ar:ApplicationResponse xmlns:ar="urn:oasis:names:specification:ubl:schema:xsd:ApplicationResponse-2"/>'),
                         ('a.pdf', b'%PDF-1.4')])
    assert len(lec.no_procesados) == 2 and not lec.comprobantes


# ----------------------------------------------------------------------------- edición
def test_cambios_en_bloque_y_restaurar(correr):                 # caso 14
    lote, ctx = correr(JOYERIA, 'comercial', archivos('joyeria_set'))
    filas = [f for f in lote.por_estado('ok') if f.c is not None]
    for f in filas:
        aplicar_cambio(f, ctx, cuenta_total='1212')
    assert {f.cuenta_total for f in filas} == {'1212'}
    for f in filas:
        f.total_manual = False
        f.cuenta_total = cuenta_total(f, ctx)
    assert {f.cuenta_total for f in filas} == {'1011'}


def test_cambiar_condicion_recalcula_cuenta_total(correr):
    lote, ctx = correr(MINERA, 'mixta', archivos('dolares_set'))
    f = fila(lote, 'E001-131')
    aplicar_cambio(f, ctx, condicion='CRE')
    assert f.cuenta_total == '1212' and f.glosa.endswith('AL CREDITO')


def test_marcar_anticipo_a_mano(correr):
    lote, ctx = correr(JOYERIA, 'comercial', archivos('joyeria_set'))
    f = fila(lote, 'BON1-521')
    aplicar_cambio(f, ctx, es_anticipo=True)
    assert f.cuenta_base == '1221' and f.anticipo_marca == 'manual'
    aplicar_cambio(f, ctx, es_anticipo=False)
    assert f.cuenta_base == '70121'


# ----------------------------------------------------------------------------- Excel
def test_excel_44_columnas_sin_cabecera(correr):
    lote, ctx = correr(JOYERIA, 'comercial', archivos('joyeria_set'))
    for f in lote.por_estado('hueco_anulado'):
        f.confirmado = True
    ws = openpyxl.load_workbook(io.BytesIO(generar_importacion(lote.filas, ctx.cfg))).active
    assert ws.max_column == 44
    assert ws.max_row == len(lote.exportables) == 33
    assert ws['A1'].number_format == 'dd/mm/yyyy'
    filas = list(ws.iter_rows(values_only=True))
    anul = next(r for r in filas if r[4] == '4480')
    assert anul[6] == '0000' and anul[7] == 'ANULADO'           # el código conserva los ceros
    assert all(r[16] for r in filas)                            # Q: TC en todas las filas
    assert all(r[24] in ('CON', 'CRE') for r in filas)


# ----------------------------------------------------------------------------- tabla editable
def test_aplicar_ediciones_de_la_tabla(correr):
    import numpy as np
    from motor.edicion import aplicar_ediciones
    lote, ctx = correr(JOYERIA, 'comercial', archivos('joyeria_set'))
    por_clave = {f.clave: f for f in lote.filas}
    a, b, c = fila(lote, 'FON1-19'), fila(lote, 'BON1-521'), fila(lote, 'BOP1-4478')
    regs = [
        {'clave': a.clave, '✓': np.bool_(True), 'Cond': 'CRE', 'AB cuenta': a.cuenta_base, 'AD cuenta': a.cuenta_total,
         'Anticipo': np.bool_(False), 'Glosa': a.glosa},
        {'clave': b.clave, '✓': False, 'Cond': b.condicion, 'AB cuenta': '70122', 'AD cuenta': b.cuenta_total,
         'Anticipo': False, 'Glosa': None},
        {'clave': c.clave, '✓': float('nan'), 'Cond': float('nan'), 'AB cuenta': c.cuenta_base, 'AD cuenta': c.cuenta_total,
         'Anticipo': float('nan'), 'Glosa': c.glosa},
    ]
    cambio, sel = aplicar_ediciones(regs, por_clave, ctx)
    assert cambio and sel == [a]
    assert (a.condicion, a.cuenta_total) == ('CRE', '1212')          # cambiar la condición recalcula AD
    assert (b.cuenta_base, b.origen, b.base_manual) == ('70122', 'editado', True)
    assert b.glosa == ''                                             # glosa vaciada a propósito
    assert (c.condicion, c.es_anticipo) == ('CON', False)            # NaN no cambia nada
    assert aplicar_ediciones(regs[2:], por_clave, ctx) == (False, [])


def test_cuenta_borrada_o_inexistente_va_a_revisar(correr):
    lote, ctx = correr(JOYERIA, 'comercial', archivos('joyeria_set'))
    f = fila(lote, 'FON1-19')
    aplicar_cambio(f, ctx, cuenta_base='')
    assert f.estado == 'revisar' and not f.exportable
    aplicar_cambio(f, ctx, cuenta_base='70122')
    assert f.estado == 'ok'
    aplicar_cambio(f, ctx, cuenta_total='9999')
    assert f.estado == 'revisar'
    aplicar_cambio(f, ctx, cuenta_total='1212')
    assert f.estado == 'ok'
